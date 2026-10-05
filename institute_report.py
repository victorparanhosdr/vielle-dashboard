"""Conservative campaign attribution and integer-cent course reporting."""
from collections import Counter, defaultdict
from datetime import timedelta
import io
import re

from institute_sources import period

PAID = {"paid", "approved"}
STATUS_LABELS = {"paid": "Aprovado", "approved": "Aprovado", "waiting_payment": "Aguardando pagamento",
    "pending": "Pendente", "processing": "Processando", "authorized": "Autorizado",
    "refunded": "Reembolsado", "chargedback": "Chargeback", "refused": "Recusado",
    "refund_requested": "Reembolso solicitado", "pending_refund": "Reembolso pendente"}


def email(value):
    value = str(value or "").strip().lower()
    return value if re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value) else ""


def phone(value):
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) in (12, 13) and digits.startswith("55"):
        digits = digits[2:]
    return digits if len(digits) in (10, 11) else ""


def identity(row):
    return "email:" + email(row.get("email")) if email(row.get("email")) else "phone:" + phone(row.get("phone")) if phone(row.get("phone")) else "id:" + row.get("course_key", "") + ":" + row["id"]


def match_lead(sale, leads, indexes=None):
    e, p = email(sale.get("email")), phone(sale.get("phone"))
    cutoff = sale.get("created_day") or sale.get("approved_day") or ""
    pool = leads if indexes is None else {r["id"]: r for r in [*indexes[0].get(e, []), *indexes[1].get(p, [])]}.values()
    candidates = [r for r in pool if r.get("created_day") and r["created_day"] <= cutoff
        and ((e and e == email(r.get("email"))) or (p and p == phone(r.get("phone"))
            and (not e or not email(r.get("email")) or e == email(r.get("email")))))]
    if not candidates:
        return None, "unmatched"
    exact = [r for r in candidates if e and email(r.get("email")) == e]
    candidates = exact or candidates
    candidates.sort(key=lambda r: (r["created_day"], r["id"]), reverse=True)
    latest = [r for r in candidates if r["created_day"] == candidates[0]["created_day"]]
    if len({r.get("campaign", "") for r in latest}) > 1:
        return None, "ambiguous"
    return candidates[0], "matched"


def paginate(rows, query):
    try:
        page = int(query.get("page", "1"))
        if not 1 <= page <= 100000:
            raise ValueError()
    except (TypeError, ValueError):
        raise ValueError("Página inválida.") from None
    search = query.get("search", "").strip().lower()
    if len(search) > 200:
        raise ValueError("Busca muito longa.")
    status = query.get("status", "")
    rows = [r for r in rows if (not search or search in " ".join(str(r.get(k, "")) for k in ("name", "email", "campaign", "id")).lower())
            and (not status or r.get("status") == status)]
    return {"rows": rows[(page-1)*50:page*50], "total": len(rows), "page": page, "pages": max(1, (len(rows)+49)//50)}


def report_sync(store, institute, course):
    if course != "all":
        return store.sync_state(institute, course)
    courses = store.courses_for(institute, course)
    result = {}
    for source in ("kiwify", "sheets", "kommo", "meta"):
        applicable = [c for c in courses if source == "kiwify" and c["product_ids"]
            or source == "sheets" and c["sheet_id"] or source == "kommo" and c["pipeline_name"]
            or source == "meta" and store.campaigns(institute, c["key"])]
        if not applicable:
            continue
        states = [store.sync_state(institute, c["key"]).get(source, {}) for c in applicable]
        errors = [c["name"] + ": " + s["error"] for c, s in zip(applicable, states) if s.get("error")]
        result[source] = {"ok": all(s.get("ok") for s in states), "count": sum(s.get("count", 0) for s in states),
            "from": max(s.get("from", "9999") for s in states), "to": min(s.get("to", "") for s in states)}
        if all(s.get("at") for s in states):
            result[source]["at"] = min(s["at"] for s in states)
        if errors:
            result[source]["error"] = " ".join(errors)
        if any(s.get("attempt_at") for s in states):
            result[source]["attempt_at"] = max(s.get("attempt_at", 0) for s in states)
    return result


def build_report(store, institute, course, query, all_rows=False):
    start, end = query.get("from", ""), query.get("to", "")
    a, b = period(start, end)
    info = store.course(institute, course)
    courses = store.courses_for(institute, course)
    if not info or not courses:
        raise LookupError()
    leads = [{**r, "course_key": c["key"], "course_name": c["name"]}
             for c in courses for r in store.records(institute, c["key"], "sheets")]
    indexes = {c["key"]: (defaultdict(list), defaultdict(list)) for c in courses}
    for lead in leads:
        if email(lead.get("email")):
            indexes[lead["course_key"]][0][email(lead["email"])].append(lead)
        if phone(lead.get("phone")):
            indexes[lead["course_key"]][1][phone(lead["phone"])].append(lead)
    # A single Kiwify order must not be counted twice if products overlap in configuration.
    sales = list({r["id"]: {**r, "course_key": c["key"], "course_name": c["name"]}
                  for c in reversed(courses) for r in store.records(institute, c["key"], "kiwify")
                  if r["product_id"] in c["product_ids"]}.values())
    mappings = [{**r, "course_key": c["key"], "course_name": c["name"]}
                for c in courses for r in store.campaigns(institute, c["key"])]
    aliases = {(r["course_key"], v): r["id"] for r in mappings for v in [r["id"], *r["aliases"]]}
    names = Counter((r["course_key"], r["name"]) for r in mappings)
    for mapping in mappings:
        token = (mapping["course_key"], mapping["name"])
        if names[token] == 1:
            aliases.setdefault(token, mapping["id"])
    selected_ids = {r["id"] for r in mappings}
    campaigns = {r["id"]: {"id": r["id"], "name": r["name"], "meta": True, "leads": 0, "sales": 0,
        "gross": 0, "net": 0, "spend": 0, "buyers": set()} for r in mappings}
    daily = {}
    d = a
    while d <= b:
        daily[d.isoformat()] = {"day": d.isoformat(), "sales": 0, "gross": 0, "net": 0, "leads": 0, "spend": 0}
        d += timedelta(days=1)
    warnings = Counter()
    pending = []

    def campaign(token, owner):
        fallback = "utm:" + (owner + ":" if course == "all" else "") + token if token else "unattributed"
        ident = aliases.get((owner, token), fallback)
        if ident not in campaigns:
            campaigns[ident] = {"id": ident, "name": token or "Sem campanha identificada (orgânico)", "meta": False,
                "leads": 0, "sales": 0, "gross": 0, "net": 0, "spend": 0, "buyers": set()}
        return campaigns[ident]

    unique = {}
    for lead in sorted(leads, key=lambda r: (r.get("created_day") or "9999", r["id"])):
        unique.setdefault(identity(lead), lead)
    cohort = {k: r for k, r in unique.items() if start <= r.get("created_day", "") <= end}
    for lead in cohort.values():
        campaign(lead.get("campaign", ""), lead["course_key"])["leads"] += 1
        daily[lead["created_day"]]["leads"] += 1
    warnings["lead_missing_date"] = sum(not r.get("created_day") for r in leads)
    approved, orders, converted, buyers, statuses = [], [], set(), set(), Counter()
    for sale in sales:
        status = sale.get("status", "unknown")
        event_day = sale.get("approved_day") if status in PAID else sale.get("refunded_day") or sale.get("created_day")
        if status in PAID and not event_day:
            if start <= sale.get("created_day", "") <= end:
                warnings["sale_missing_date"] += 1
                pending.append({"id": sale["id"], "name": sale.get("name"), "reason": "Pagamento aprovado sem data de aprovação"})
            continue
        if not event_day or not start <= event_day <= end:
            continue
        statuses[status] += 1
        lead, association = match_lead(sale, leads, indexes[sale["course_key"]])
        token = sale.get("campaign") or (lead or {}).get("campaign") or ""
        bucket = campaign(token, sale["course_key"])
        orders.append({**sale, "status_label": STATUS_LABELS.get(status, status), "day": event_day,
                       "campaign": bucket["name"], "attribution": "UTM Kiwify" if sale.get("campaign") else "Formulário" if lead and token else "Sem atribuição"})
        if status not in PAID:
            continue
        if sale.get("currency") != "BRL" or sale.get("net_currency") != "BRL" or sale.get("gross") is None or sale.get("net") is None:
            warnings["sale_invalid_amount"] += 1
            pending.append({"id": sale["id"], "name": sale.get("name"), "reason": "Moeda diferente de BRL ou valores incompletos"})
            continue
        approved.append(sale)
        buyer = identity(sale)
        buyers.add(buyer)
        bucket["sales"] += 1
        bucket["gross"] += sale["gross"]
        bucket["net"] += sale["net"]
        bucket["buyers"].add(buyer)
        daily[event_day]["sales"] += 1
        daily[event_day]["gross"] += sale["gross"]
        daily[event_day]["net"] += sale["net"]
        if lead and identity(lead) in cohort:
            converted.add(identity(lead))
        if not token or association == "ambiguous" and not sale.get("campaign"):
            warnings["sale_attribution"] += 1
            pending.append({"id": sale["id"], "name": sale.get("name"), "reason": "Campanha ausente ou vínculo ambíguo"})
    seen_meta = set()
    for c in courses:
        owned_ids = {r["id"] for r in mappings if r["course_key"] == c["key"]}
        for row in store.records(institute, c["key"], "meta"):
            ident = (row["campaign_id"], row["day"])
            if row["campaign_id"] in owned_ids and start <= row["day"] <= end and ident not in seen_meta:
                seen_meta.add(ident)
                campaigns[row["campaign_id"]]["spend"] += row["spend"]
                daily[row["day"]]["spend"] += row["spend"]
    sync = report_sync(store, institute, course)
    meta_state = sync.get("meta", {})
    covered = bool(selected_ids) and meta_state.get("ok") and meta_state.get("from", "9999") <= start and meta_state.get("to", "") >= end
    config = store.public_settings(institute)
    meta_status = "ready" if covered else "connection_missing" if not (config.get("meta_access_token_configured") and config.get("meta_account_id")) else "campaigns_missing" if not selected_ids else "sync_failed" if meta_state.get("error") and meta_state.get("attempt_at") else "period_missing"
    for row in campaigns.values():
        spend = row["spend"]
        valid = row["meta"] and covered
        row.update({"spend": spend if valid else None,
            "cpl": round(spend / row["leads"]) if valid and row["leads"] else None,
            "cac": round(spend / len(row["buyers"])) if valid and row["buyers"] else None,
            "roas": row["gross"] / spend if valid and spend else None})
        del row["buyers"]
    paid_campaigns = [r for r in campaigns.values() if r["meta"]]
    spend = sum(r["spend"] or 0 for r in paid_campaigns) if covered else None
    gross, net = sum(r["gross"] for r in approved), sum(r["net"] for r in approved)
    crm = [{**r, "course_key": c["key"], "course_name": c["name"]} for c in courses
           for r in store.records(institute, c["key"], "kommo") if c["pipeline_name"]
           and r["pipeline_name"].strip().casefold() == c["pipeline_name"].strip().casefold()
           and start <= r.get("created_day", "") <= end]
    stages = {}
    for r in crm:
        stage = (r["course_key"], r["stage"])
        name = r["course_name"] + " · " + r["stage"] if course == "all" and sum(bool(c["pipeline_name"]) for c in courses) > 1 else r["stage"]
        item = stages.setdefault(stage, {"name": name, "sort": r["stage_sort"], "count": 0})
        item["count"] += 1
    order_page = paginate(sorted(orders, key=lambda r: (r["day"], r["id"]), reverse=True), query)
    if all_rows:
        order_page["rows"] = [r for r in sorted(orders, key=lambda r: (r["day"], r["id"]), reverse=True)
            if (not query.get("search") or query["search"].strip().lower() in " ".join(str(r.get(k, "")) for k in ("name", "email", "campaign", "id")).lower())
            and (not query.get("status") or r["status"] == query["status"])]
    return {"ok": True, "course": info, "period": {"from": start, "to": end}, "sync": sync,
        "summary": {"sales": len(approved), "gross": gross, "net": net, "leads": len(cohort), "buyers": len(buyers),
            "conversion": len(converted)/len(cohort)*100 if cohort else None, "converted_leads": len(converted),
            "ticket": round(gross/len(approved)) if approved else None, "spend": spend,
            "roas": sum(r["gross"] for r in paid_campaigns)/spend if spend else None,
            "net_after_ads": net-spend if spend is not None else None},
        "daily": list(daily.values()), "campaigns": sorted(campaigns.values(), key=lambda r: (-r["gross"], r["name"])),
        "stages": sorted(stages.values(), key=lambda r: r["sort"]),
        "statuses": [{"key": k, "name": STATUS_LABELS.get(k, k), "count": v} for k, v in statuses.items()],
        "orders": order_page,
        "leads": paginate(sorted(cohort.values(), key=lambda r: (r["created_day"], r["id"]), reverse=True), query),
        "warnings": {k: v for k, v in warnings.items() if v}, "pending": pending[:200],
        "totals_partial": bool(warnings["sale_missing_date"] or warnings["sale_invalid_amount"]), "meta_coverage": bool(covered),
        "meta_status": meta_status, "sources": [s for s in ("kiwify", "sheets", "kommo", "meta")
            if s == "meta" or s == "kiwify" and any(c["product_ids"] for c in courses)
            or s == "sheets" and any(c["sheet_id"] for c in courses) or s == "kommo" and any(c["pipeline_name"] for c in courses)]}


def workbook(store, institute, course, query):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    report = build_report(store, institute, course, query, all_rows=True)
    book = Workbook()
    summary = book.active
    summary.title = "Resumo"
    summary.append(["Curso", report["course"]["name"]])
    summary.append(["De", query["from"], "Até", query["to"]])
    for field, label in (("sales", "Vendas aprovadas"), ("leads", "Leads únicos"), ("gross", "Receita bruta"), ("net", "Receita líquida"), ("spend", "Investimento Meta")):
        value = report["summary"][field]
        summary.append([label, value/100 if value is not None and field in {"gross", "net", "spend"} else value])
    summary.append(["Base", "Aprovação Kiwify; líquido não significa saldo disponível para saque."])
    sheets = [("Campanhas", ["Campanha", "Leads", "Vendas", "Bruto", "Líquido", "Investimento", "CPL", "CAC", "ROAS"],
        [[r["name"], r["leads"], r["sales"], *[r[k]/100 if r[k] is not None else None for k in ("gross", "net", "spend", "cpl", "cac")], r["roas"]] for r in report["campaigns"]]),
        ("Dia a dia", ["Data", "Leads", "Vendas", "Bruto", "Líquido", "Investimento"],
        [[r["day"], r["leads"], r["sales"], r["gross"]/100, r["net"]/100, r["spend"]/100 if report["meta_coverage"] else None] for r in report["daily"]])]
    all_orders = report["orders"]["rows"]
    sheets.append(("Pedidos", ["ID", "Data", "Aluno", "E-mail", "Status", "Campanha", "Bruto", "Líquido", "Moeda", "Curso"],
        [[r["id"], r["day"], r["name"], r["email"], r["status_label"], r["campaign"],
          r["gross"]/100 if r["gross"] is not None else None, r["net"]/100 if r["net"] is not None else None, r["currency"], r["course_name"]] for r in all_orders]))
    for title, headers, rows in sheets:
        ws = book.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for ws in book:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="123B30")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(50, max(15, max(len(str(c.value or "")) for c in col)+2))
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
