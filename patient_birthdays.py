"""Clinic-scoped birthdays and annual gift records; never writes to Experts."""
import calendar
from datetime import date, datetime
from io import BytesIO
import json
import logging
import math
import re
import sqlite3
from urllib.parse import parse_qs, quote
from zoneinfo import ZoneInfo

from financial_receipts import cancelled, normalized, record


class Conflict(ValueError):
    pass


def today():
    return datetime.now(ZoneInfo("America/Sao_Paulo")).date()


def initialize(conn):
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS birthday_gifts (
            patient_uuid TEXT NOT NULL, year INTEGER NOT NULL,
            sent_at TEXT, description TEXT NOT NULL, note TEXT NOT NULL,
            actor_id INTEGER NOT NULL, actor_name TEXT NOT NULL,
            updated_at TEXT NOT NULL, revision INTEGER NOT NULL,
            PRIMARY KEY (patient_uuid, year)
        );
        CREATE TABLE IF NOT EXISTS birthday_gift_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            patient_uuid TEXT NOT NULL, year INTEGER NOT NULL,
            action TEXT NOT NULL, actor_id INTEGER NOT NULL,
            actor_name TEXT NOT NULL, at TEXT NOT NULL, details TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS birthday_gift_history
            ON birthday_gift_events(patient_uuid, id DESC);
    """)


def birth_date(raw, reference):
    for key in ("date_birth", "birth_date", "date_of_birth", "birthdate", "birthday", "data_nascimento"):
        value = raw.get(key)
        if not value:
            continue
        if not isinstance(value, str):
            return None
        try:
            result = date.fromisoformat(value[:10]) if re.match(r"^\d{4}-\d{2}-\d{2}", value) else datetime.strptime(value, "%d/%m/%Y").date()
            return result if date(1900, 1, 1) <= result <= reference else None
        except ValueError:
            return None
    return None


def removed_patient(raw):
    return raw.get("deleted") is True or bool(raw.get("deleted_at")) or normalized(raw.get("status")) in {"deleted", "removed", "excluido", "excluida"}


def options(query, reference=None):
    reference = reference or today()
    def one(key, default=""):
        values = query.get(key, [default])
        if len(values) != 1:
            raise ValueError("Informe apenas um valor por filtro.")
        return values[0]
    month = one("month", reference.strftime("%Y-%m"))
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ValueError("Mês inválido.")
    year, number = map(int, month.split("-"))
    if not 2000 <= year <= reference.year + 1 or not 1 <= number <= 12:
        raise ValueError("Mês inválido.")
    status, purchases, sort = one("status", "all"), one("purchases", "all"), one("sort", "birthday")
    if status not in {"all", "sent", "pending"} or purchases not in {"all", "with", "without"} or sort not in {"birthday", "amount", "sales", "name"}:
        raise ValueError("Filtro inválido.")
    q = one("q").strip()
    if len(q) > 200:
        raise ValueError("Busca muito longa.")
    try:
        page = int(one("page", "1"))
        if not 1 <= page <= 100000:
            raise ValueError()
    except ValueError:
        raise ValueError("Página inválida.") from None
    return {"month": month, "year": year, "number": number, "status": status, "purchases": purchases, "sort": sort, "q": q, "page": page}


def patient(conn, uuid):
    row = conn.execute("SELECT uuid,name,phone,email,raw_json FROM clinica_patients WHERE uuid=?", (uuid,)).fetchone()
    if not row or removed_patient(record(row["raw_json"])):
        raise LookupError("Paciente não encontrado nesta clínica.")
    return row


def purchase_history(conn, uuids):
    from app import money_value, sale_status_group
    totals = {uuid: {"sales": 0, "amount": 0, "last_sale": None, "missing_amount": 0, "items": []} for uuid in uuids}
    all_items = []
    # Only UUID-linked sales participate. Names and phone numbers are never matching keys.
    for start in range(0, len(uuids), 400):
        keys = uuids[start:start + 400]
        marks = ",".join("?" for _ in keys)
        for row in conn.execute(f"SELECT uuid,patient_uuid,type,sale_date,total,raw_json FROM clinica_sales WHERE patient_uuid IN ({marks})", keys):
            raw = record(row["raw_json"])
            if cancelled(raw) or raw.get("active") is False or sale_status_group(raw.get("status"), raw.get("type") or row["type"]) != "venda":
                continue
            identities = {str(value.get("uuid") or value.get("id")) for key in ("buyer", "patient", "person", "buyer_person")
                          if isinstance((value := raw.get(key)), dict) and (value.get("uuid") or value.get("id"))}
            if identities and identities != {row["patient_uuid"]}:
                continue
            amount = money_value(raw.get("final_amount")) if raw.get("final_amount") is not None else row["total"]
            try:
                amount = float(amount) if amount is not None else None
                if amount is not None and (not math.isfinite(amount) or amount < 0):
                    amount = None
            except (ValueError, TypeError):
                amount = None
            day = str(row["sale_date"] or "")[:10]
            try:
                day = date.fromisoformat(day).isoformat()
            except ValueError:
                day = None
            value = totals[row["patient_uuid"]]
            value["sales"] += 1
            if amount is None:
                value["missing_amount"] += 1
            else:
                value["amount"] += round(amount * 100)
            if day and day > (value["last_sale"] or ""):
                value["last_sale"] = day
            item = {"uuid": row["uuid"], "date": day, "amount": amount}
            value["items"].append(item)
            all_items.append(item)
    for value in totals.values():
        value["amount"] /= 100
        value["items"].sort(key=lambda r: (r["date"] or "", r["uuid"]), reverse=True)
    return totals, all_items


def report(conn, filters, reference=None, export=False):
    reference = reference or today()
    rows, missing, synced = [], 0, 0
    for row in conn.execute("SELECT uuid,name,phone,email,raw_json,synced_at FROM clinica_patients"):
        raw = record(row["raw_json"])
        if removed_patient(raw):
            continue
        synced = max(synced, row["synced_at"] or 0)
        birthday = birth_date(raw, reference)
        if not birthday:
            missing += 1
            continue
        if birthday.month != filters["number"] or birthday.year > filters["year"]:
            continue
        observed_day = min(birthday.day, calendar.monthrange(filters["year"], birthday.month)[1])
        occurrence = date(filters["year"], birthday.month, observed_day)
        rows.append({"uuid": row["uuid"], "name": row["name"] or "Paciente sem nome", "phone": row["phone"] or "", "email": row["email"] or "",
                     "birthday": occurrence.isoformat(), "birth_day": birthday.day, "age": filters["year"] - birthday.year,
                     "is_today": occurrence == reference, "shifted": observed_day != birthday.day})
    amounts, _ = purchase_history(conn, [r["uuid"] for r in rows])
    gifts = {r["patient_uuid"]: dict(r) for r in conn.execute("SELECT * FROM birthday_gifts WHERE year=?", (filters["year"],))}
    distribution = [0, 0, 0, 0]
    daily = [0] * calendar.monthrange(filters["year"], filters["number"])[1]
    for row in rows:
        values = amounts[row["uuid"]]
        row.update({k: v for k, v in values.items() if k != "items"})
        row["gift"] = gifts.get(row["uuid"])
        row["gift_sent"] = bool(row["gift"] and row["gift"]["sent_at"])
        distribution[0 if not row["sales"] else 1 if row["sales"] == 1 else 2 if row["sales"] <= 4 else 3] += 1
        daily[int(row["birthday"][-2:]) - 1] += 1
    sent = sum(r["gift_sent"] for r in rows)
    totals = {"patients": len(rows), "with_purchases": sum(r["sales"] > 0 for r in rows), "sent": sent, "pending": len(rows) - sent,
              "sales": sum(r["sales"] for r in rows), "amount": round(sum(r["amount"] for r in rows), 2), "today": sum(r["is_today"] for r in rows)}
    warnings = {"missing_birth_dates": missing, "missing_sale_amounts": sum(r["missing_amount"] for r in rows)}
    query = normalized(filters["q"])
    selected = [r for r in rows if (not query or query in normalized(" ".join((r["name"], r["phone"], r["email"]))))
                and (filters["status"] == "all" or r["gift_sent"] == (filters["status"] == "sent"))
                and (filters["purchases"] == "all" or bool(r["sales"]) == (filters["purchases"] == "with"))]
    key = {"birthday": lambda r: (r["birthday"], normalized(r["name"])), "amount": lambda r: (-r["amount"], r["birthday"], normalized(r["name"])),
           "sales": lambda r: (-r["sales"], r["birthday"], normalized(r["name"])), "name": lambda r: (normalized(r["name"]), r["birthday"])}[filters["sort"]]
    selected.sort(key=key)
    count = len(selected)
    if not export:
        selected = selected[(filters["page"] - 1) * 30:filters["page"] * 30]
    return {"ok": True, "month": filters["month"], "year": filters["year"], "today": reference.isoformat(), "totals": totals, "warnings": warnings,
            "synced_at": synced, "daily": daily, "distribution": distribution, "patients": selected, "filtered_count": count,
            "page": filters["page"], "pages": max(1, (count + 29) // 30), "basis": "Compras acumuladas na base sincronizada do Clínica Experts."}


def detail(conn, uuid):
    row = patient(conn, uuid)
    purchases, _ = purchase_history(conn, [uuid])
    history = [dict(r) for r in conn.execute("SELECT * FROM birthday_gifts WHERE patient_uuid=? ORDER BY year DESC", (uuid,))]
    events = [dict(r) for r in conn.execute("SELECT id,year,action,actor_name,at,details FROM birthday_gift_events WHERE patient_uuid=? ORDER BY id DESC LIMIT 30", (uuid,))]
    for event in events:
        event["details"] = json.loads(event["details"])
    return {"ok": True, "patient": {"uuid": uuid, "name": row["name"] or "Paciente sem nome"}, "history": history, "events": events,
            "purchases": {**purchases[uuid], "items": purchases[uuid]["items"][:20]}}


def text(value, limit, required=False):
    if not isinstance(value, str) or len(value.strip()) > limit or any(ord(c) < 32 and c not in "\n\t" for c in value):
        raise ValueError("Texto inválido ou muito longo.")
    value = value.strip()
    if required and not value:
        raise ValueError("Informe o presente enviado.")
    return value


def save_gift(conn, payload, actor, reference=None, undo=False):
    reference = reference or today()
    uuid = text(payload.get("patient_uuid"), 128, True)
    patient(conn, uuid)
    year, revision = payload.get("year"), payload.get("revision")
    if type(year) is not int or not 2000 <= year <= reference.year or type(revision) is not int or revision < 0:
        raise ValueError("Ano ou revisão inválidos.")
    old = conn.execute("SELECT * FROM birthday_gifts WHERE patient_uuid=? AND year=?", (uuid, year)).fetchone()
    if revision != (old["revision"] if old else 0):
        raise Conflict("Este registro foi alterado por outra pessoa. Reabra para conferir.")
    if undo:
        if not old or not old["sent_at"]:
            raise ValueError("Não há envio para desfazer.")
        sent_at, description, note = None, old["description"], text(payload.get("note", ""), 2000, True)
    else:
        try:
            sent_at = date.fromisoformat(payload.get("sent_at", ""))
        except (ValueError, TypeError):
            raise ValueError("Informe a data do envio.") from None
        if sent_at.year != year or sent_at > reference:
            raise ValueError("A data deve pertencer ao ano selecionado e não pode ser futura.")
        sent_at = sent_at.isoformat()
        description, note = text(payload.get("description"), 200, True), text(payload.get("note", ""), 2000)
    stamp = datetime.now(ZoneInfo("America/Sao_Paulo")).isoformat(timespec="seconds")
    values = (sent_at, description, note, actor["id"], actor["nome"], stamp, revision + 1, uuid, year)
    if old:
        cursor = conn.execute("""UPDATE birthday_gifts SET sent_at=?,description=?,note=?,actor_id=?,actor_name=?,updated_at=?,revision=?
                                 WHERE patient_uuid=? AND year=? AND revision=?""", (*values, revision))
        if cursor.rowcount != 1:
            raise Conflict("Este registro mudou. Reabra para conferir.")
    else:
        try:
            conn.execute("""INSERT INTO birthday_gifts(sent_at,description,note,actor_id,actor_name,updated_at,revision,patient_uuid,year)
                            VALUES(?,?,?,?,?,?,?,?,?)""", values)
        except sqlite3.IntegrityError:
            raise Conflict("Este registro mudou. Reabra para conferir.") from None
    conn.execute("INSERT INTO birthday_gift_events(patient_uuid,year,action,actor_id,actor_name,at,details) VALUES(?,?,?,?,?,?,?)",
                 (uuid, year, "undone" if undo else "updated" if old else "sent", actor["id"], actor["nome"], stamp,
                  json.dumps({"before": dict(old) if old else None, "sent_at": sent_at, "description": description, "note": note}, ensure_ascii=False)))
    return detail(conn, uuid)


def workbook(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = "Aniversariantes"
    sheet.append(["Paciente", "Aniversário", "Idade no aniversário", "Telefone", "E-mail", "Vendas", "Total comprado (R$)", "Última compra",
                  "Ano do presente", "Presente enviado", "Data do envio", "Presente", "Observação", "Registrado por"])
    for row in data["patients"]:
        gift = row["gift"] or {}
        sheet.append([row["name"], row["birthday"], row["age"], row["phone"], row["email"], row["sales"], row["amount"], row["last_sale"],
                      data["year"], "Sim" if row["gift_sent"] else "Não", gift.get("sent_at"), gift.get("description"), gift.get("note"), gift.get("actor_name")])
    for row in sheet:
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="153D31")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in "ABCDEFGHJKLMN":
        sheet.column_dimensions[column].width = 23
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["M"].width = 40
    for row in sheet.iter_rows(min_row=2):
        row[6].number_format = '#,##0.00'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    criteria = book.create_sheet("Critérios")
    for row in [("Mês", data["month"]), ("Base de compras", data["basis"]), ("Pacientes sem data de nascimento", data["warnings"]["missing_birth_dates"]),
                ("Vendas sem valor confirmado", data["warnings"]["missing_sale_amounts"]), ("Vínculo", "Identificador do paciente, sem associação por nome."),
                ("Exclusões", "Orçamentos, vendas canceladas/excluídas e pacientes excluídos."), ("29 de fevereiro", "Em anos não bissextos, exibido em 28 de fevereiro.")]:
        criteria.append(row)
    criteria.column_dimensions["A"].width = 32
    criteria.column_dimensions["B"].width = 90
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def handle_request(handler, parsed, connect):
    query = parse_qs(parsed.query, keep_blank_values=True)
    clinic = query["clinic"][0]
    try:
        with connect() as conn:
            initialize(conn)
            if handler.command == "GET":
                if parsed.path == "/api/birthdays/patient":
                    if len(query.get("id", [])) != 1:
                        raise ValueError("Informe o paciente.")
                    return handler.auth_json(detail(conn, query["id"][0]))
                if parsed.path not in {"/api/birthdays", "/api/birthdays/export"}:
                    return handler.auth_json({"ok": False, "error": "Rota não encontrada."}, 404)
                data = report(conn, options(query), export=parsed.path.endswith("/export"))
                if parsed.path.endswith("/export"):
                    content = workbook(data)
                    handler.send_response(200)
                    handler.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    handler.send_header("Content-Length", str(len(content)))
                    handler.send_header("Content-Disposition", "attachment; filename*=UTF-8''" + quote("aniversariantes-" + clinic + "-" + data["month"] + ".xlsx"))
                    handler.end_headers()
                    handler.wfile.write(content)
                    return
                return handler.auth_json(data)
            if parsed.path not in {"/api/birthdays/gift", "/api/birthdays/gift/undo"}:
                return handler.auth_json({"ok": False, "error": "Rota não encontrada."}, 404)
            length = int(handler.headers.get("Content-Length", "0"))
            if handler.headers.get_content_type() != "application/json" or handler.headers.get("Transfer-Encoding") or not 0 < length <= 12000:
                raise ValueError("Solicitação inválida.")
            payload = json.loads(handler.rfile.read(length))
            if not isinstance(payload, dict):
                raise ValueError("Solicitação inválida.")
            if not isinstance(payload.get("patient_uuid"), str) or type(payload.get("year")) is not int:
                raise ValueError("Paciente ou ano inválidos.")
            old = conn.execute("SELECT revision FROM birthday_gifts WHERE patient_uuid=? AND year=?", (payload.get("patient_uuid"), payload.get("year"))).fetchone()
            if not handler.require_permission(clinic, "birthdays.edit" if old or parsed.path.endswith("/undo") else "birthdays.create"):
                return
            data = save_gift(conn, payload, handler.current_user, undo=parsed.path.endswith("/undo"))
        return handler.auth_json(data)
    except Conflict as error:
        return handler.auth_json({"ok": False, "error": str(error)}, 409)
    except (ValueError, TypeError, OverflowError, UnicodeError) as error:
        return handler.auth_json({"ok": False, "error": str(error) if isinstance(error, ValueError) else "Dados inválidos."}, 400)
    except LookupError as error:
        return handler.auth_json({"ok": False, "error": str(error)}, 404)
    except sqlite3.OperationalError:
        return handler.auth_json({"ok": False, "error": "Banco ocupado. Tente novamente."}, 503)
    except Exception:
        logging.exception("Birthday operation failed for clinic %s", clinic)
        return handler.auth_json({"ok": False, "error": "Não foi possível concluir a operação."}, 500)
