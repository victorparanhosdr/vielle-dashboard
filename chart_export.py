"""Excel exports from the same filtered snapshot used by the dashboard."""
import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CHARTS = {
    "revenue_daily": "Evolução do faturamento diário",
    "sales_ticket": "Vendas e ticket médio por dia",
    "accumulated": "Faturamento acumulado",
    "top_patients": "Top pacientes",
    "value_ranges": "Distribuição por faixa de valor",
    "leads_bookings": "Leads x agendamentos criados por dia",
    "expense_categories": "Saídas por categoria",
    "expense_daily": "Saídas por dia",
}

LABELS = {
    "uuid": "ID de origem", "id": "ID Kommo", "patient_uuid": "ID paciente",
    "patient": "Paciente", "professional": "Profissional", "professional_uuid": "ID profissional",
    "procedure_uuid": "ID procedimento", "bill_uuid": "ID lançamento",
    "sale_date": "Data da venda", "date": "Data considerada", "day": "Dia considerado",
    "registered_at": "Criado em", "starts_at": "Início da consulta", "ends_at": "Fim da consulta",
    "due_date": "Vencimento", "paid_at": "Pagamento", "emission_date": "Emissão",
    "amount": "Valor considerado (R$)", "total": "Total de origem (R$)",
    "settled": "Pago/recebido (R$)", "open_amount": "Em aberto (R$)",
    "description": "Descrição", "detail": "Categoria considerada", "category_name": "Categoria original",
    "account_name": "Conta", "type": "Tipo", "status": "Status", "direction": "Movimento",
    "name": "Nome", "price": "Valor (R$)", "pipeline_name": "Funil", "status_name": "Etapa",
    "responsible_user_id": "ID responsável", "registry_user_name": "Registrado por",
    "source": "Fonte", "created_at": "Criação (timestamp Kommo)",
    "updated_at": "Atualização (timestamp Kommo)", "closed_at": "Fechamento (timestamp Kommo)",
}
MONEY = {"amount", "total", "settled", "open_amount", "price"}


def add_sheet(wb, name, headers, rows, money_columns=()):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        values = [json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in row]
        if any(isinstance(v, str) and len(v) > 32767 for v in values):
            raise ValueError("Um campo excede o limite do Excel.")
        if ws.max_row >= 1048576:
            raise ValueError("O período excede o limite de linhas do Excel. Reduza o período.")
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                # Treat CRM content as text, never as executable Excel formulas.
                cell.data_type = "s"
            if cell.column in money_columns:
                cell.number_format = '"R$" #,##0.00'
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="123B30")
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(20, len(header) + 3))
    return ws


def details_sheet(wb, name, records, originals):
    keys = list(dict.fromkeys(k for row in records for k in row if k != "raw_json"))
    keys = keys or ["uuid", "date", "amount"]
    rows = []
    for record in records:
        rows.append([str(record[k]) if record.get(k) is not None and (k == "id" or k.endswith("_id"))
                     else record.get(k) for k in keys])
        raw = record.get("raw_json")
        if raw:
            # Original API record is lossless, even beyond Excel's per-cell limit.
            raw = raw if isinstance(raw, str) else json.dumps(raw, ensure_ascii=False)
            for offset in range(0, len(raw), 30000):
                originals.append([name, str(record.get("uuid", record.get("id", ""))),
                                  offset // 30000 + 1, raw[offset:offset + 30000]])
    add_sheet(wb, name, [LABELS.get(k, k) for k in keys], rows,
              [i + 1 for i, key in enumerate(keys) if key in MONEY])


def build_workbook(chart, panel, context):
    if chart not in CHARTS:
        raise ValueError("Gráfico inválido.")
    wb = Workbook()
    wb.remove(wb.active)
    details = panel.get("export_details", {})
    originals = []
    financial = panel.get("financial_daily", [])
    if chart in {"revenue_daily", "accumulated", "expense_daily"}:
        cumulative = 0
        rows = []
        for item in financial:
            amount = item.get("expenses" if chart == "expense_daily" else "income", 0)
            cumulative += amount
            rows.append([item["day"], amount, cumulative])
        add_sheet(wb, "Resumo", ["Dia", "Valor do dia (R$)", "Acumulado no período (R$)"], rows, [2, 3])
    elif chart == "sales_ticket":
        add_sheet(wb, "Resumo", ["Dia", "Vendas", "Faturamento (R$)", "Ticket médio (R$)"],
                  [[r["day"], r["sales"], r["revenue"], r["average_ticket"]]
                   for r in panel.get("sales_ticket_daily", [])], [3, 4])
    elif chart == "value_ranges":
        add_sheet(wb, "Resumo", ["Faixa", "Vendas", "Valor (R$)", "Participação (0 a 1)"],
                  [[r["range"], r["sales"], r["amount"], r["share"]]
                   for r in panel.get("value_ranges", [])], [3])
    elif chart == "top_patients":
        patients = {}
        for row in details.get("sales", []):
            key = row.get("patient_uuid") or row.get("patient")
            item = patients.setdefault(key, [str(key or ""), row.get("patient"), 0, 0])
            item[2] += 1
            item[3] += row.get("amount", 0)
        add_sheet(wb, "Ranking completo", ["ID paciente", "Paciente", "Vendas", "Faturamento (R$)"],
                  sorted(patients.values(), key=lambda r: r[3], reverse=True), [4])
    elif chart == "expense_categories":
        categories = panel.get("expenses_by_category", [])
        total = sum(r["amount"] for r in categories)
        add_sheet(wb, "Resumo", ["Categoria", "Valor (R$)", "Participação (0 a 1)"],
                  [[r["category"], r["amount"], r["amount"] / total if total else 0] for r in categories], [2])
    else:
        leads = {r["day"]: r.get("total", 0) for r in panel.get("daily_leads", [])}
        bookings = {r["day"]: r.get("total", 0) for r in panel.get("daily_bookings", [])}
        add_sheet(wb, "Resumo", ["Dia de criação", "Leads", "Agendamentos incluídos"],
                  [[day, leads.get(day, 0), bookings.get(day, 0)] for day in sorted(leads.keys() | bookings.keys())])
    sources = ({"Leads": "leads", "Agendamentos": "bookings"} if chart == "leads_bookings" else
               {"Saídas": "expenses"} if chart.startswith("expense_") else
               {"Vendas": "sales"} if chart in {"sales_ticket", "top_patients", "value_ranges"} else
               {"Lançamentos": "income"})
    for name, key in sources.items():
        details_sheet(wb, name, details.get(key, []), originals)
    add_sheet(wb, "Registros originais", ["Aba", "ID origem", "Parte JSON", "JSON original (concatenar partes)"], originals)
    add_sheet(wb, "Contexto", ["Campo", "Valor"], [
        ["Gráfico", CHARTS[chart]], *context.items(),
        ["Gerado em", datetime.now().isoformat(timespec="seconds")],
        ["Origem", "Base sincronizada do sistema; mesmos filtros e regras do gráfico."],
        ["Detalhamento", "Todos os registros considerados, sem limite de linhas da interface."],
        ["Data de referência", "Leads: criação Kommo. Agendamentos: registered_at ou created_at; se ausentes, o gráfico atual usa starts_at. Confira Dia considerado e as datas originais."
         if chart == "leads_bookings" else "Vendas: data da venda. Financeiro: pagamento, vencimento ou emissão, conforme o gráfico."],
    ])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
