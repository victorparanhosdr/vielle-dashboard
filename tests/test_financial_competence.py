import io
import json
import sqlite3
import unittest

from openpyxl import load_workbook
from financial_competence import build_report, query_options, export_workbook
from financial_receipts import build_receipts


class CompetenceTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.addCleanup(self.conn.close)
        self.conn.executescript("""
            create table clinica_bills(uuid text, raw_json text, synced_at int, type text, emission_date text);
            create table clinica_parcels(uuid text, bill_uuid text, status text, raw_json text, synced_at int, type text, paid_at text);
            create table clinica_sales(patient_uuid text, sale_date text, type text, raw_json text);
        """)
        self.options = query_options({"date_from": ["2026-09-01"], "date_to": ["2026-09-30"]})

    def add(self, uuid="bill", kind="Venda", amount=10000, net=9700, day="2026-09-10", status="open", **extra):
        parcel = {"uuid": "parcel-" + uuid, "status": status, "due_date": "2026-10-10", "compensation_date": "2026-11-10"}
        raw = {"uuid": uuid, "type": kind, "final_amount": amount, "net_amount": net, "emission_date": day,
               "description": "Título fictício " + uuid, "category": {"name": "Serviços"},
               "person": {"uuid": "person-" + uuid, "name": "Contato fictício " + uuid},
               "payment_methods": [{"parcels": [parcel]}], **extra}
        self.conn.execute("insert into clinica_bills(uuid,raw_json,synced_at) values(?,?,10)", (uuid, json.dumps(raw)))
        self.conn.execute("insert into clinica_parcels(uuid,bill_uuid,status,raw_json,synced_at) values(?,?,?,?,11)",
                          (parcel["uuid"], uuid, status, json.dumps(parcel)))
        return raw

    def report(self, **changes):
        return build_report(self.conn, {**self.options, **changes})

    def payments(self, raw, parcels):
        raw["payment_methods"] = [{"parcels": parcels}]
        self.conn.execute("update clinica_bills set raw_json=?,synced_at=20 where uuid=?", (json.dumps(raw), raw["uuid"]))
        self.conn.execute("delete from clinica_parcels where bill_uuid=?", (raw["uuid"],))
        for parcel in parcels:
            self.conn.execute("insert into clinica_parcels(uuid,bill_uuid,status,raw_json,synced_at) values(?,?,?,?,20)",
                              (parcel["uuid"], raw["uuid"], parcel["status"], json.dumps(parcel)))

    def test_partial_payment_splits_actual_net_and_outstanding_gross(self):
        raw = self.add(net=9800, seller={"uuid": "doctor-a"})
        self.payments(raw, [
            {"uuid": "received", "status": "received", "final_amount": 3000, "net_amount": 2900, "compensation_date": "2026-09-15"},
            {"uuid": "open", "status": "open", "final_amount": 7000, "net_amount": 6900, "due_date": "2026-10-15"},
        ])
        r = self.report()
        item = r["items"][0]
        self.assertEqual((item["gross"], item["net"], item["received"], item["open"], item["fees"], item["payment_state"]),
                         (100, 98, 29, 70, 2, "partial"))
        self.assertEqual((r["settlement_totals"]["income_received"], r["settlement_totals"]["income_open"]), (29, 70))
        self.assertEqual(self.report(payment_status="not_received")["count"], 0)
        self.assertEqual(self.report(payment_status="received")["count"], 0)
        self.assertEqual(self.report(payment_status="partial")["count"], 1)
        cash = self.report(date_basis="receipt")
        self.assertEqual((cash["count"], cash["totals"]["income"], cash["totals"]["income_gross"]), (1, 29, 30))
        self.assertEqual(cash["items"][0]["fees"], 1)
        self.assertEqual(cash["totals"]["income"], build_receipts(self.conn, "2026-09-01", "2026-09-30")["net_total"])

    def test_receipt_basis_matches_card_including_old_titles_manual_and_all_professionals(self):
        for key, kind, amount, doctor in (("sale", "Venda", 10000, "doctor-a"), ("manual", "Conta", 5000, None),
                                         ("custom", "Honorários contábeis", 2000, "doctor-b")):
            raw = self.add(uuid=key, kind=kind, amount=amount, net=amount, day="2026-08-10", seller={"uuid": doctor})
            self.payments(raw, [{"uuid": "p-" + key, "status": "received", "final_amount": amount,
                                 "net_amount": amount - 100, "compensation_date": "2026-09-15"}])
        self.assertEqual(self.report()["count"], 0)
        for doctors in ([], ["doctor-a"], ["doctor-b"]):
            with self.subTest(doctors=doctors):
                r = build_report(self.conn, {**self.options, "date_basis": "receipt"}, doctors)
                card = build_receipts(self.conn, "2026-09-01", "2026-09-30", doctors)
                self.assertEqual((r["totals"]["income"], r["count"]), (card["net_total"], card["count"]))
                self.assertEqual(r["settlement_totals"]["income_received"], card["net_total"])
        self.assertEqual(self.report(date_basis="receipt")["totals"]["income"], 167)

    def test_status_filters_distinguish_full_partial_and_unpaid(self):
        raw = self.add(uuid="received")
        self.payments(raw, [{"uuid": "r", "status": "received", "final_amount": 10000, "net_amount": 9700, "paid_at": "2026-09-12"}])
        raw = self.add(uuid="open")
        self.payments(raw, [{"uuid": "o", "status": "late", "final_amount": 10000, "net_amount": 9700}])
        self.assertEqual(self.report(payment_status="received")["items"][0]["payment_state"], "received")
        self.assertEqual(self.report(payment_status="not_received")["settlement_totals"]["income_open"], 100)
        self.assertEqual(self.report(payment_status="not_received")["settlement_totals"]["income_received"], 0)
        self.assertEqual(self.report(date_basis="receipt", payment_status="not_received")["count"], 0)

    def test_missing_dates_partial_balance_and_future_compensation_never_invent_receipt(self):
        for key, extra in (("missing", {"status": "received", "due_date": "2026-09-10", "calc_compensation_date": "2026-09-10"}),
                           ("future", {"status": "received", "compensation_date": "2099-09-10"}),
                           ("partial", {"status": "partial", "balance": 5000})):
            raw = self.add(uuid=key)
            self.payments(raw, [{"uuid": "p-" + key, "final_amount": 10000, "net_amount": 9700, **extra}])
        r = self.report()
        self.assertEqual(r["settlement_totals"]["income_received"], 0)
        self.assertEqual(r["settlement_totals"]["income_open"], 100)
        self.assertEqual({item["id"] for item in r["payment_pending"]}, {"missing", "partial"})
        self.assertEqual(self.report(date_basis="receipt")["totals"]["income"], 0)

    def test_received_date_filters_each_parcel_and_deduplicates_flat_snapshot(self):
        raw = self.add(amount=20000, net=19400)
        self.payments(raw, [
            {"uuid": "a", "status": "received", "final_amount": 10000, "net_amount": 9700, "compensation_date": "2026-08-30"},
            {"uuid": "b", "status": "received", "final_amount": 10000, "net_amount": 9700, "compensation_date": "2026-09-12"},
        ])
        r = self.report(date_basis="receipt", date_from="2026-09-12", date_to="2026-09-12")
        self.assertEqual((r["count"], r["totals"]["income"]), (1, 97))
        self.assertEqual(r["items"][0]["uuid"], "b")
        self.assertEqual(self.report()["settlement_totals"]["income_received"], 194)
        self.conn.execute("update clinica_parcels set status='open',synced_at=21 where uuid='b'")
        self.assertEqual(self.report(date_basis="receipt")["totals"]["income"], 0)

    def test_expenses_payment_basis_is_net_paid_not_open_or_title_date(self):
        raw = self.add(kind="Conta", day="2026-08-10", seller={"uuid": "doctor-a"}, amount=20000, net=19400)
        self.payments(raw, [
            {"uuid": "p", "status": "paid", "final_amount": 10000, "net_amount": 9700, "payment_date": "2026-09-12"},
            {"uuid": "o", "status": "open", "final_amount": 10000, "net_amount": 9700, "due_date": "2026-09-12"},
        ])
        r = self.report(date_basis="receipt")
        self.assertEqual((r["totals"]["income"], r["totals"]["expense"], r["settlement_totals"]["expense_paid"]), (0, 97, 97))
        self.assertEqual(r["items"][0]["net"], -97)
        self.assertEqual(r["items"][0]["payment_state"], "partial")

    def test_receipt_with_missing_title_date_or_gross_still_matches_card(self):
        raw = self.add(emission_date=None)
        self.payments(raw, [{"uuid": "p", "status": "received", "net_amount": 9700, "compensation_date": "2026-09-15"}])
        r = self.report(date_basis="receipt")
        self.assertEqual(r["totals"]["income"], 97)
        self.assertIsNone(r["items"][0]["gross"])
        self.assertEqual(r["gross_incomplete"], 1)

    def test_cancelled_titles_and_parcels_do_not_count_as_paid(self):
        raw = self.add()
        self.payments(raw, [{"uuid": "p", "status": "received", "final_amount": 10000, "net_amount": 9700, "compensation_date": "2026-09-15"}])
        self.conn.execute("update clinica_bills set raw_json=json_set(raw_json,'$.deleted',json('true'))")
        self.assertEqual(self.report()["count"], 0)
        self.assertEqual(self.report(date_basis="receipt")["totals"]["income"], 0)

    def test_export_uses_current_basis_status_and_settlement_columns(self):
        raw = self.add()
        self.payments(raw, [{"uuid": "p", "status": "received", "final_amount": 10000, "net_amount": 9700, "compensation_date": "2026-09-15"}])
        options = {**self.options, "date_basis": "receipt", "payment_status": "received"}
        report = build_report(self.conn, options, export=True)
        workbook = load_workbook(io.BytesIO(export_workbook(report, options, "Teste")))
        self.assertEqual(workbook["Recebimentos"]["N2"].value, 3)
        self.assertEqual(workbook["Recebimentos"]["O2"].value, 97)
        self.assertEqual(workbook["Recebimentos"]["R2"].value, "bill")
        for params in ({"date_basis": ["due"]}, {"payment_status": ["bad"]}, {"payment_status": ["all", "received"]}):
            with self.assertRaises(ValueError):
                query_options(params)

    def test_date_is_competence_not_payment_due_or_creation(self):
        self.add(created_at="2026-12-15")
        self.assertEqual(self.report(date_from="2026-09-10", date_to="2026-09-10")["totals"]["income"], 97)
        self.assertEqual(self.report(date_from="2026-10-01", date_to="2026-11-30")["count"], 0)
        self.add(uuid="explicit", competence_date="2026-09-11")
        self.assertEqual(self.report(date_from="2026-09-11", date_to="2026-09-11")["items"][0]["date_source"], "competence_date")

    def test_missing_emission_never_falls_back_to_due_or_created(self):
        self.add(emission_date=None, created_at="2026-09-10", due_date="2026-09-10")
        self.assertEqual(self.report()["count"], 0)
        self.assertEqual(self.report()["excluded"]["date"], 1)

    def test_title_once_not_each_installment_and_include_unpaid(self):
        raw = self.add()
        raw["payment_methods"][0]["parcels"].append({"uuid": "second", "status": "late"})
        self.conn.execute("update clinica_bills set raw_json=?", (json.dumps(raw),))
        self.conn.execute("insert into clinica_parcels(uuid,bill_uuid,status,raw_json,synced_at) values('second','bill','late','{}',10)")
        self.assertEqual((self.report()["count"], self.report()["totals"]["income"]), (1, 97))

    def test_income_expense_manual_receipt_and_negative_net(self):
        self.add()
        self.add(uuid="expense", kind="Conta", amount=6000, net=6000, status="paid")
        self.add(uuid="manual", kind="Conta", amount=4000, net=4000, status="received")
        report = self.report(direction="expense")
        self.assertEqual(report["totals"], {"income": 137, "expense": 60, "balance": 77,
                                           "income_gross": 140, "expense_gross": 60, "balance_gross": 80})
        self.assertEqual(report["count"], 1)
        self.assertEqual(report["items"][0]["net"], -60)
        self.assertEqual(report["items"][0]["gross"], 60)

    def test_open_expenses_and_manual_receivable(self):
        self.add(kind="Conta", status="open")
        self.add(uuid="income", kind="Conta", status="open", description="Título a receber de teste")
        self.assertEqual(self.report()["totals"], {"income": 97, "expense": 97, "balance": 0,
                                                 "income_gross": 100, "expense_gross": 100, "balance_gross": 0})

    def test_gross_and_net_totals_follow_filters_before_direction_and_pagination(self):
        for number in range(61):
            self.add(uuid=str(number), amount=10001, net=9701, seller={"uuid": "doctor-a"})
        self.add(uuid="expense", kind="Conta", amount=6001, net=5501, seller={"uuid": "doctor-a"})
        self.add(uuid="other-doctor", amount=90000, net=80000, seller={"uuid": "doctor-b"})
        self.add(uuid="other-month", day="2026-08-10", seller={"uuid": "doctor-a"})
        report = build_report(self.conn, {**self.options, "direction": "income", "page": 2}, ["doctor-a"])
        self.assertEqual(len(report["items"]), 11)
        self.assertEqual(report["totals"], {"income": 5917.61, "expense": 55.01, "balance": 5862.60,
                                           "income_gross": 6100.61, "expense_gross": 60.01, "balance_gross": 6040.60})
        contact = build_report(self.conn, {**self.options, "contact": "person-expense"}, ["doctor-a"])
        self.assertEqual(contact["totals"], {"income": 0, "expense": 55.01, "balance": -55.01,
                                            "income_gross": 0, "expense_gross": 60.01, "balance_gross": -60.01})
        self.assertEqual(self.report(search="inexistente")["totals"], dict.fromkeys(report["totals"], 0))

    def test_cancelled_initial_balance_and_unknown_direction(self):
        self.add(status="cancelled")
        self.add(uuid="balance", kind="Saldo Inicial")
        self.add(uuid="unknown", kind="Transferência")
        self.assertEqual(self.report()["count"], 0)
        self.assertEqual(self.report()["excluded"]["direction"], 1)

    def test_latest_parcel_status_controls_manual_receipt(self):
        self.add(kind="Conta", status="received")
        self.conn.execute("update clinica_parcels set status='paid'")
        self.assertEqual(self.report()["items"][0]["direction"], "expense")
        self.conn.execute("update clinica_bills set synced_at=12")
        self.assertEqual(self.report()["items"][0]["direction"], "income")

    def test_embedded_orphan_parent_deduplicated_and_missing_parent_reported(self):
        raw = self.add()
        self.conn.execute("delete from clinica_bills")
        self.conn.execute("update clinica_parcels set raw_json=?", (json.dumps({"raw_bill": raw}),))
        self.assertEqual(self.report()["count"], 1)
        self.conn.execute("insert into clinica_parcels(uuid,bill_uuid,status,raw_json,synced_at) values('orphan','unknown','paid','{}',1)")
        self.assertEqual(self.report()["excluded"]["date"], 1)

    def test_missing_net_and_explicit_fee_fallback_zero_retained(self):
        self.add(net=None)
        self.add(uuid="fees", net=None, fees_amount=300)
        self.add(uuid="zero", net=0)
        report = self.report()
        self.assertEqual(report["totals"]["income"], 97)
        self.assertEqual(report["excluded"]["amount"], 1)
        self.assertEqual(report["count"], 2)

    def test_filters_accent_insensitive_search_and_category_type_contact(self):
        self.add(description="Avaliação clínica", category={"name": "Consulta"})
        self.add(uuid="other", kind="Conta", category={"name": "Aluguel"})
        self.assertEqual(self.report(search="avaliacao")["count"], 1)
        self.assertEqual(self.report(contact="person-bill")["count"], 1)
        self.assertEqual(self.report(category="Aluguel", title_type="Venda")["count"], 0)
        self.assertEqual(len(self.report(category="Aluguel")["options"]["categories"]), 2)

    def test_professional_matching_is_exact_and_does_not_mix(self):
        self.add()
        self.add(uuid="direct", seller={"uuid": "doctor-b"})
        self.add(uuid="missing")
        self.conn.execute("insert into clinica_sales values('person-bill','2026-09-10','sale',?)",
                          (json.dumps({"final_amount": 10000, "seller": {"uuid": "doctor-a"}}),))
        a = build_report(self.conn, self.options, ["doctor-a"])
        b = build_report(self.conn, self.options, ["doctor-b"])
        self.assertEqual([row["uuid"] for row in a["items"]], ["bill"])
        self.assertEqual([row["uuid"] for row in b["items"]], ["direct"])
        self.assertEqual(a["excluded"]["professional"], 1)
        self.conn.execute("insert into clinica_sales values('person-bill','2026-09-10','sale',?)",
                          (json.dumps({"final_amount": 10000, "seller": {"uuid": "doctor-b"}}),))
        self.assertEqual(build_report(self.conn, self.options, ["doctor-a"])["count"], 0)

    def test_export_all_pages_totals_and_formula_safety(self):
        for number in range(61):
            self.add(uuid=str(number), description="=HYPERLINK(\"bad\")")
        self.assertEqual(len(self.report()["items"]), 50)
        self.assertEqual(len(self.report(page=2)["items"]), 11)
        self.assertEqual(self.report(page=999)["page"], 2)
        report = build_report(self.conn, self.options, export=True)
        workbook = load_workbook(io.BytesIO(export_workbook(report, self.options, "Clínica teste")))
        self.assertEqual(workbook["Competência"].max_row, 62)
        self.assertEqual(workbook["Competência"]["B2"].data_type, "s")
        self.assertEqual(sum(row[7] for row in workbook["Competência"].iter_rows(min_row=2, values_only=True)), 61 * 97)
        self.assertEqual(workbook["Resumo"]["B2"].value, 61 * 100)
        self.assertEqual(workbook["Resumo"]["C2"].value, 61 * 97)
        self.assertEqual(workbook["Resumo"]["B4"].value, report["totals"]["balance_gross"])
        self.assertEqual(workbook["Resumo"]["C4"].value, report["totals"]["balance"])

    def test_invalid_dates_filters_sort_page_and_duplicate_parameters(self):
        for params in ({"date_from": ["2026-09-31"]}, {"date_from": ["2026-10-02"], "date_to": ["2026-10-01"]},
                       {"sort": ["raw_json"]}, {"direction": ["bad"]}, {"page": ["0"]},
                       {"category": ["a", "b"]}, {"search": ["a" * 501]}):
            with self.assertRaises(ValueError):
                query_options(params)

    def test_pending_identify_problem_records_without_counting_them(self):
        self.add(uuid="valid")
        self.add(uuid="no-date", emission_date=None)
        self.add(uuid="no-amount", net=None)
        self.add(uuid="unknown-direction", kind="Transferência")
        self.add(uuid="cancelled", status="cancelled")
        report = self.report()
        self.assertEqual(report["totals"]["income"], 97)
        self.assertEqual({item["id"]: item["reason"] for item in report["pending"]},
                         {"no-date": "date", "no-amount": "amount", "unknown-direction": "direction"})
        self.assertEqual(len(report["pending"]), sum(report["excluded"].values()))
        self.assertEqual(report["pending"][0]["contact"], "Contato fictício no-date")

    def test_pending_professional_and_dates_respect_report_scope(self):
        self.add(uuid="unknown", seller=None)
        self.add(uuid="other", net=None, seller={"uuid": "doctor-b"})
        self.add(uuid="other-undated", emission_date=None, seller={"uuid": "doctor-b"})
        self.add(uuid="last-month", day="2026-08-10")
        report = build_report(self.conn, self.options, ["doctor-a"])
        self.assertEqual([(item["id"], item["reason"]) for item in report["pending"]], [("unknown", "professional")])


if __name__ == "__main__":
    unittest.main()
