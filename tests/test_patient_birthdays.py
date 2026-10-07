import http.client
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
from http.server import ThreadingHTTPServer
from unittest.mock import patch

import app
from auth_store import AuthStore
import patient_birthdays as birthdays


DAY = date(2026, 10, 7)
ACTOR = {"id": 7, "nome": "Recepção"}


def seed(conn, uuid="patient", birth="1985-10-07", name="Ana Costa", **extra):
    raw = {"uuid": uuid, "date_birth": birth, **extra}
    conn.execute("INSERT INTO clinica_patients(uuid,name,phone,email,raw_json,synced_at) VALUES(?,?,?,?,?,?)",
                 (uuid, name, "5531999990000", uuid + "@example.invalid", json.dumps(raw), 100))


def sale(conn, uuid, buyer="patient", kind="sale", status="active", amount=12500, **extra):
    raw = {"type": kind, "status": status, "final_amount": amount, "buyer": {"uuid": buyer}, **extra}
    conn.execute("INSERT INTO clinica_sales(uuid,patient_uuid,type,sale_date,total,raw_json,synced_at) VALUES(?,?,?,?,?,?,0)",
                 (uuid, buyer, kind, extra.get("sale_date", "2025-06-01"), 9999, json.dumps(raw)))


class BirthdayTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "clinic.sqlite3"
        patcher = patch.object(app, "clinic_db_path", return_value=self.path)
        patcher.start(); self.addCleanup(patcher.stop)
        app.init_db()
        self.conn = app.db(); self.addCleanup(self.conn.close)
        birthdays.initialize(self.conn)
        seed(self.conn)

    def report(self, **query):
        return birthdays.report(self.conn, birthdays.options({k: [str(v)] for k, v in query.items()}, DAY), DAY)

    def gift(self, **extra):
        return birthdays.save_gift(self.conn, {"patient_uuid": "patient", "year": 2026, "revision": 0,
            "sent_at": "2026-10-07", "description": "Kit de cuidados", "note": "Entregue", **extra}, ACTOR, DAY)

    def test_current_month_source_fields_and_no_guessed_birthdays(self):
        seed(self.conn, "previous", "1980-09-07")
        seed(self.conn, "missing", None)
        seed(self.conn, "invalid", "1980-02-30")
        seed(self.conn, "future", "2027-10-07")
        seed(self.conn, "deleted", "1980-10-07", deleted=True)
        seed(self.conn, "br", "08/10/1990", name="Bia")
        data = self.report()
        self.assertEqual(data["month"], "2026-10")
        self.assertEqual(data["totals"]["patients"], 2)
        self.assertEqual(data["patients"][0]["age"], 41)
        self.assertTrue(data["patients"][0]["is_today"])
        self.assertEqual(data["warnings"]["missing_birth_dates"], 3)
        self.assertEqual(sum(data["daily"]), 2)

    def test_linked_sales_use_final_amount_and_exclude_quotes_cancelled_conflicts(self):
        sale(self.conn, "one", amount=12500)
        sale(self.conn, "zero", amount=0)
        sale(self.conn, "quote", kind="sale_quote", status="won")
        sale(self.conn, "cancelled", status="cancelled")
        sale(self.conn, "removed", deleted_at="2026-01-01")
        sale(self.conn, "inactive", status="inactive")
        sale(self.conn, "conflicting", patient={"uuid": "other"})
        sale(self.conn, "unlinked", buyer="other")
        data = self.report()
        row = data["patients"][0]
        self.assertEqual(row["sales"], 2)
        self.assertEqual(row["amount"], 125)
        self.assertEqual(row["last_sale"], "2025-06-01")
        self.assertEqual(data["distribution"], [0, 0, 1, 0])
        self.assertEqual(data["totals"]["with_purchases"], 1)

    def test_annual_gifts_edit_undo_and_history_survive_reopening(self):
        self.gift()
        self.assertEqual(self.report()["totals"]["sent"], 1)
        self.assertEqual(self.report(month="2025-10")["totals"]["sent"], 0)
        self.gift(revision=1, note="Recebido pela paciente")
        with self.assertRaises(birthdays.Conflict): self.gift(revision=1)
        birthdays.save_gift(self.conn, {"patient_uuid": "patient", "year": 2026, "revision": 2, "note": "Registro incorreto"}, ACTOR, DAY, undo=True)
        self.assertEqual(self.report()["totals"]["pending"], 1)
        self.conn.commit()
        with sqlite3.connect(self.path) as conn:
            conn.row_factory = sqlite3.Row
            data = birthdays.detail(conn, "patient")
            self.assertEqual(len(data["history"]), 1)
            self.assertEqual(data["history"][0]["revision"], 3)
            self.assertEqual([e["action"] for e in data["events"]], ["undone", "updated", "sent"])
            self.assertEqual(data["events"][0]["actor_name"], "Recepção")

    def test_purchase_history_includes_2024_2025_and_2026(self):
        for year in (2024, 2025, 2026):
            sale(self.conn, str(year), amount=10000, sale_date=f"{year}-01-05")
        row = self.report()["patients"][0]
        self.assertEqual((row["sales"], row["amount"], row["last_sale"]), (3, 300, "2026-01-05"))

    def test_header_sort_directions_apply_before_pagination(self):
        for n in range(40):
            seed(self.conn, f"p{n}", "1990-10-15", name=f"Paciente {n:02}")
            sale(self.conn, f"s{n}", buyer=f"p{n}", amount=n * 100, sale_date=f"2025-06-{n % 28 + 1:02}")
        self.assertEqual(self.report(sort="amount")["patients"][0]["uuid"], "p39")
        descending = self.report(sort="amount", direction="desc")["patients"]
        ascending = self.report(sort="amount", direction="asc")["patients"]
        self.assertGreater(descending[0]["amount"], descending[-1]["amount"])
        self.assertLess(ascending[0]["amount"], ascending[-1]["amount"])
        page2 = self.report(sort="amount", direction="desc", page=2)["patients"]
        self.assertGreaterEqual(descending[-1]["amount"], page2[0]["amount"])
        self.assertEqual(self.report(sort="name", direction="desc")["patients"][0]["name"], "Paciente 39")
        self.assertEqual(self.report(sort="last_sale", direction="asc", page=2)["patients"][-1]["uuid"], "patient")
        self.gift()
        self.assertEqual(self.report(sort="gift", direction="desc")["patients"][0]["uuid"], "patient")

    def test_filters_sort_pagination_and_safe_filtered_workbook(self):
        for n in range(40):
            seed(self.conn, f"p{n}", "1990-10-15", name=f"Paciente {n:02}")
        seed(self.conn, "formula", "1990-10-16", name="=DANGEROUS()")
        self.gift()
        self.assertEqual(len(self.report()["patients"]), 30)
        self.assertEqual(len(self.report(page=2)["patients"]), 12)
        self.assertEqual(self.report(status="sent")["filtered_count"], 1)
        self.assertEqual(self.report(status="pending")["filtered_count"], 41)
        self.assertEqual(self.report(q="ana")["filtered_count"], 1)
        sale(self.conn, "one", buyer="p0")
        self.assertEqual(self.report(purchases="with", sort="amount")["patients"][0]["uuid"], "p0")
        options = birthdays.options({"status": ["pending"]}, DAY)
        data = birthdays.report(self.conn, options, DAY, export=True)
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(birthdays.workbook(data)))
        self.assertEqual(book.active.max_row, 42)
        self.assertTrue(all(c.data_type != "f" for r in book.active for c in r))
        self.assertEqual(sum(data["distribution"]), data["totals"]["patients"])

    def test_leap_birthdays_and_each_year_status(self):
        seed(self.conn, "leap", "2000-02-29")
        normal = self.report(month="2026-02")["patients"][0]
        self.assertEqual(normal["birthday"], "2026-02-28")
        self.assertTrue(normal["shifted"])
        leap = self.report(month="2024-02")["patients"][0]
        self.assertEqual(leap["birthday"], "2024-02-29")
        self.assertFalse(leap["shifted"])

    def test_input_validation(self):
        for extra in ({"sent_at": "2026-10-08"}, {"sent_at": "2025-10-07"}, {"year": True}, {"revision": True},
                      {"description": ""}, {"description": "x" * 201}, {"patient_uuid": "unknown"}):
            with self.assertRaises((ValueError, LookupError)): self.gift(**extra)
        for query in ({"month": ["2026-13"]}, {"month": ["2026-10", "2026-09"]}, {"sort": ["raw_json"]}, {"direction": ["random"]}, {"page": ["0"]}):
            with self.assertRaises(ValueError): birthdays.options(query, DAY)


class BirthdayHttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(); cls.root = Path(cls.temp.name)
        cls.patcher = patch.object(app, "clinic_db_path", side_effect=lambda: cls.root / (app.current_clinic_id() + ".sqlite3"))
        cls.patcher.start()
        for clinic in ("vielle", "inspire"):
            with app.clinic_context(clinic), app.db() as conn:
                seed(conn, "shared", name="Paciente " + clinic)
        cls.auth = AuthStore(cls.root / "auth.sqlite3"); cls.auth.initialize()
        cls.auth.create_first_master("Master", "master", "Birthday-tests-password!")
        for name, permissions in (("writer", ["view", "create", "edit", "export"]), ("reader", ["view"]), ("none", [])):
            cls.auth.create_user(name, name, "Birthday-tests-password!", clinic_keys=["vielle"], permissions={"vielle": ["birthdays." + p for p in permissions]})
        cls.tokens = {name: cls.auth.login(name, "Birthday-tests-password!") for name in ("master", "writer", "reader", "none")}
        class Quiet(app.Handler):
            def log_message(self, *_): pass
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), Quiet)
        cls.server.auth_store, cls.server.login_limiter = cls.auth, app.LoginLimiter()
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True); cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join(); cls.patcher.stop(); cls.temp.cleanup()

    def req(self, path, user="writer", payload=None, csrf=True, method=None):
        headers = {"Cookie": "doc4docs_session=" + self.tokens[user]} if user else {}
        if csrf: headers["X-DOC4DOCS-Request"] = "1"
        if payload is not None: headers["Content-Type"] = "application/json"
        conn = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=10)
        conn.request(method or ("POST" if payload is not None else "GET"), path, json.dumps(payload) if payload is not None else None, headers)
        response = conn.getresponse(); data = response.read(); status = response.status; kind = response.getheader("Content-Type", ""); conn.close()
        return status, json.loads(data) if kind.startswith("application/json") and data else data

    def test_access_permissions_aliases_and_csrf(self):
        self.assertEqual(self.req("/api/birthdays?clinic=vielle", user=None)[0], 401)
        self.assertEqual(self.req("/birthdays.html?clinic=vielle", user=None)[0], 303)
        for path in ("/birthdays.html?clinic=vielle", "/static/birthdays.html?clinic=vielle", "/%62irthdays.html?clinic=vielle", "/api/birthdays?clinic=vielle"):
            self.assertEqual(self.req(path, user="none")[0], 403)
        self.assertEqual(self.req("/api/birthdays?clinic=inspire")[0], 403)
        self.assertEqual(self.req("/api/birthdays?clinic=vielle&clinic=inspire")[0], 400)
        self.assertEqual(self.req("/api/birthdays")[0], 400)
        self.assertEqual(self.req("/api/birthdays/export?clinic=vielle", user="reader")[0], 403)
        self.assertEqual(self.req("/api/birthdays/export?clinic=vielle", method="POST", payload={})[0], 405)
        self.assertEqual(self.req("/api/birthdays/gift?clinic=vielle", payload={}, csrf=False)[0], 403)
        self.assertEqual(self.req("/api/birthdays?clinic=vielle", method="HEAD")[0], 405)

    def test_saved_gift_stays_in_clinic_and_reader_cannot_mutate(self):
        payload = {"patient_uuid": "shared", "year": 2026, "revision": 0, "sent_at": "2026-10-07", "description": "Kit", "note": ""}
        with patch.object(birthdays, "today", return_value=DAY):
            self.assertEqual(self.req("/api/birthdays/gift?clinic=vielle", user="reader", payload=payload)[0], 403)
            self.assertEqual(self.req("/api/birthdays/gift?clinic=vielle", payload=payload)[0], 200)
            self.assertEqual(self.req("/api/birthdays/gift?clinic=vielle", payload=payload)[0], 409)
            self.assertEqual(self.req("/api/birthdays?clinic=vielle")[1]["totals"]["sent"], 1)
            self.assertEqual(self.req("/api/birthdays?clinic=inspire", user="master")[1]["totals"]["sent"], 0)
            self.assertEqual(self.req("/api/birthdays/patient?clinic=inspire&id=shared", user="master")[1]["history"], [])
            self.assertEqual(self.req("/api/birthdays/gift?clinic=vielle", payload={"patient_uuid": [], "year": 2026})[0], 400)

    def test_history_access_connection_and_csrf(self):
        from birthday_history import JOBS
        self.assertEqual(self.req("/api/birthdays/history?clinic=vielle", user="none")[0], 403)
        self.assertEqual(self.req("/api/birthdays/history?clinic=inspire")[0], 403)
        self.assertEqual(self.req("/api/birthdays/history?clinic=vielle", payload={}, csrf=False)[0], 403)
        with patch.object(app, "config_value", return_value=""):
            self.assertEqual(self.req("/api/birthdays/history?clinic=vielle")[0], 200)
            self.assertEqual(self.req("/api/birthdays/history?clinic=vielle", payload={})[0], 409)
        with patch.object(app, "config_value", return_value="connected"), patch("birthday_history.start", return_value={"ok": True}) as start:
            self.assertEqual(self.req("/api/birthdays/history?clinic=vielle", user="reader", payload={})[0], 202)
            start.assert_called_once_with("vielle")
        self.assertFalse(JOBS.status("vielle")["running"])


if __name__ == "__main__": unittest.main()
