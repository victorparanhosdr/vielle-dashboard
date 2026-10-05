import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from institute_store import InstituteStore, INSTITUTES
from institute_sources import Kiwify, Meta, SourceError, day, sheet_rows, windows, kommo
from institute_report import build_report, match_lead, workbook
from institute_api import sync_sources

I, C = "victor-paranhos", "regen-code"


class InstitutesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = InstituteStore(Path(self.tmp.name) / "institutes.sqlite3")
        self.store.initialize()
        c = self.store.course(I, C)
        c.pop("institute_key")
        c["product_ids"] = ["regen"]
        self.store.save_course(I, c)

    def lead(self, ident="a", **extra):
        return {"id": ident, "name": "Aluno teste", "email": "aluno@example.invalid", "phone": "+5511999990000",
                "created_day": "2026-09-01", "campaign": "co2", **extra}

    def sale(self, ident="s", **extra):
        return {"id": ident, "name": "Aluno teste", "email": "aluno@example.invalid", "phone": "11999990000", "product_id": "regen",
            "created_day": "2026-09-05", "approved_day": "2026-09-06", "status": "paid", "gross": 100000, "net": 90000,
            "currency": "BRL", "net_currency": "BRL", "campaign": "", "payment_method": "pix", **extra}

    def report(self, **options):
        return build_report(self.store, I, C, {"from": "2026-09-01", "to": "2026-09-30", **options})

    def test_members_not_clinic_memberships(self):
        self.assertEqual(self.store.allowed({"id": 4, "is_master": False}), [])
        self.store.save_members(I, [4])
        self.assertEqual(self.store.allowed({"id": 4, "is_master": False}), [I])
        self.assertEqual(self.store.allowed({"id": 9, "is_master": True}), list(INSTITUTES))

    def test_presencial_course_seed_is_additive_and_idempotent(self):
        course = self.store.course(I, "regen-code-presencial")
        self.assertEqual(len(course["product_ids"]), 2)
        self.assertEqual((course["sheet_id"], course["pipeline_name"]), ("", ""))
        course.pop("institute_key")
        course["name"] = "Presencial personalizado"
        self.store.save_course(I, course)
        self.store.save_records(I, C, "kiwify", [self.sale()])
        self.store.initialize()
        self.assertEqual(self.store.course(I, "regen-code-presencial")["name"], "Presencial personalizado")
        self.assertEqual(self.store.course(I, C)["product_ids"], ["regen"])
        self.assertEqual(len(self.store.records(I, C, "kiwify")), 1)
        self.assertEqual(len(self.store.courses(I)), 2)

    def test_presencial_reports_do_not_mix_online_sales(self):
        course = "regen-code-presencial"
        product = self.store.course(I, course)["product_ids"][0]
        self.store.save_records(I, C, "kiwify", [self.sale()])
        self.store.save_records(I, course, "kiwify", [self.sale("presencial", product_id=product), self.sale("online")])
        report = build_report(self.store, I, course, {"from": "2026-09-01", "to": "2026-09-30"})
        self.assertEqual(report["summary"]["sales"], 1)
        self.assertEqual(report["orders"]["rows"][0]["id"], "presencial")
        self.assertEqual(self.report()["orders"]["rows"][0]["id"], "s")

    def test_presencial_sync_skips_unconfigured_leads_and_crm(self):
        with patch("institute_api.Kiwify") as kiwify, patch("institute_api.sheets") as sheet, patch("institute_api.kommo") as crm:
            kiwify.return_value.sales.return_value = []
            results = sync_sources(self.store, I, "regen-code-presencial", "2026-09-01", "2026-09-30", None, lambda _: None)
        self.assertEqual(set(results), {"kiwify"})
        self.assertTrue(results["kiwify"]["ok"])
        sheet.assert_not_called()
        crm.assert_not_called()

    def test_all_courses_combines_products_without_duplicate_orders(self):
        presencial = "regen-code-presencial"
        info = self.store.course(I, presencial)
        product = info["product_ids"][0]
        info.pop("institute_key")
        info["product_ids"].append("regen")
        self.store.save_course(I, info)
        self.store.save_records(I, C, "kiwify", [self.sale(), self.sale("clinic", product_id="clinic")])
        self.store.save_records(I, presencial, "kiwify", [self.sale(), self.sale("presencial", product_id=product, gross=200000, net=180000)])
        report = build_report(self.store, I, "all", {"from":"2026-09-01", "to":"2026-09-30"})
        self.assertEqual((report["summary"]["sales"], report["summary"]["gross"], report["summary"]["net"]), (2, 300000, 270000))
        self.assertEqual(report["summary"]["buyers"], 1)
        self.assertEqual({r["id"] for r in report["orders"]["rows"]}, {"s", "presencial"})
        self.assertTrue(report["course"]["is_all"])
        self.assertEqual(sum(r["gross"] for r in report["daily"]), 300000)

    def test_all_courses_lead_matching_stays_with_its_course(self):
        presencial = "regen-code-presencial"
        product = self.store.course(I, presencial)["product_ids"][0]
        self.store.save_records(I, C, "sheets", [self.lead()])
        self.store.save_records(I, C, "kiwify", [self.sale()])
        self.store.save_records(I, presencial, "kiwify", [self.sale("presencial", product_id=product)])
        report = build_report(self.store, I, "all", {"from":"2026-09-01", "to":"2026-09-30"})
        order = next(r for r in report["orders"]["rows"] if r["id"] == "presencial")
        self.assertEqual(order["attribution"], "Sem atribuição")
        self.assertEqual(report["summary"]["leads"], 1)
        self.assertEqual(report["summary"]["conversion"], 100)
        self.store.save_records(I, presencial, "sheets", [self.lead("another", created_day="2026-09-02")])
        self.assertEqual(build_report(self.store, I, "all", {"from":"2026-09-01", "to":"2026-09-30"})["summary"]["leads"], 1)

    def test_all_courses_meta_sums_only_owned_campaigns_and_requires_coverage(self):
        presencial = "regen-code-presencial"
        self.store.save_settings(I, {"meta_access_token":"test", "meta_account_id":"123456"})
        for owner, campaign, spend in ((C, "123", 10000), (presencial, "456", 20000)):
            self.store.save_campaigns(I, owner, [{"id":campaign, "name":owner, "aliases":[]}])
            self.store.save_records(I, owner, "meta", [{"id":campaign+":2026-09-01", "campaign_id":campaign, "day":"2026-09-01", "spend":spend},
                {"id":"999:2026-09-01", "campaign_id":"999", "day":"2026-09-01", "spend":999999}])
            self.store.set_sync_state(I, owner, "meta", {"ok":True, "at":1, "from":"2026-09-01", "to":"2026-09-30"})
        report = build_report(self.store, I, "all", {"from":"2026-09-01", "to":"2026-09-30"})
        self.assertEqual(report["summary"]["spend"], 30000)
        self.assertEqual(sum(r["spend"] for r in report["daily"]), 30000)
        self.assertEqual(report["meta_status"], "ready")
        self.store.set_sync_state(I, presencial, "meta", {"ok":True, "at":1, "from":"2026-09-02", "to":"2026-09-30"})
        report = build_report(self.store, I, "all", {"from":"2026-09-01", "to":"2026-09-30"})
        self.assertIsNone(report["summary"]["spend"])
        self.assertEqual(report["meta_status"], "period_missing")

    def test_meta_pending_distinguishes_configuration_and_refresh(self):
        self.assertEqual(self.report()["meta_status"], "connection_missing")
        self.store.save_settings(I, {"meta_access_token":"test", "meta_account_id":"123456"})
        self.assertEqual(self.report()["meta_status"], "campaigns_missing")
        self.store.save_campaigns(I, C, [{"id":"123", "name":"co2", "aliases":[]}])
        self.store.set_sync_state(I, C, "meta", {"ok":False, "error":"Campanhas alteradas. Atualize o investimento para este período."})
        self.assertEqual(self.report()["meta_status"], "period_missing")
        self.store.set_sync_state(I, C, "meta", {"ok":False, "error":"Meta indisponível", "attempt_at":1})
        self.assertEqual(self.report()["meta_status"], "sync_failed")

    def test_exact_meta_campaign_name_attributes_without_manual_alias(self):
        self.store.save_campaigns(I, C, [{"id":"123", "name":"co2", "aliases":[]}])
        self.store.save_records(I, C, "kiwify", [self.sale(campaign="co2")])
        campaign = next(r for r in self.report()["campaigns"] if r["sales"])
        self.assertEqual(campaign["id"], "123")
        self.assertTrue(campaign["meta"])

    def test_duplicate_meta_campaign_names_are_not_guessed(self):
        self.store.save_campaigns(I, C, [{"id":"123", "name":"co2", "aliases":[]}, {"id":"456", "name":"co2", "aliases":[]}])
        self.store.save_records(I, C, "kiwify", [self.sale(campaign="co2")])
        campaign = next(r for r in self.report()["campaigns"] if r["sales"])
        self.assertFalse(campaign["meta"])

    def test_all_courses_sync_updates_each_real_course(self):
        with patch("institute_api.Kiwify") as kiwify, patch("institute_api.sheets", return_value=[]) as sheet, patch("institute_api.kommo", return_value=[]) as crm:
            kiwify.return_value.sales.return_value = []
            result = sync_sources(self.store, I, "all", "2026-09-01", "2026-09-30", None, lambda _: None)
        self.assertEqual(kiwify.return_value.sales.call_count, 2)
        self.assertEqual({call.args[0]["key"] for call in kiwify.return_value.sales.call_args_list}, {C, "regen-code-presencial"})
        sheet.assert_called_once()
        crm.assert_called_once()
        self.assertTrue(result["kiwify"]["ok"])

    def test_all_filter_cannot_be_saved_as_a_course(self):
        with self.assertRaises(ValueError):
            self.store.save_course(I, {"key":"all", "name":"Todos", "product_ids":[]})

    def test_all_courses_export_includes_course_and_combined_total(self):
        from openpyxl import load_workbook
        presencial = "regen-code-presencial"
        product = self.store.course(I, presencial)["product_ids"][0]
        self.store.save_records(I, C, "kiwify", [self.sale()])
        self.store.save_records(I, presencial, "kiwify", [self.sale("presencial", product_id=product)])
        book = load_workbook(io.BytesIO(workbook(self.store, I, "all", {"from":"2026-09-01", "to":"2026-09-30"})))
        self.assertEqual(book["Pedidos"].max_row, 3)
        self.assertEqual(book["Pedidos"]["J1"].value, "Curso")
        self.assertEqual(book["Resumo"]["B1"].value, "Todos os cursos")

    def test_secrets_not_returned_and_blank_preserves(self):
        self.store.save_settings(I, {"kiwify_client_secret": "very-private", "meta_access_token": "private-meta"})
        self.store.save_settings(I, {"kiwify_client_secret": ""})
        public = self.store.public_settings(I)
        self.assertNotIn("very-private", json.dumps(public))
        self.assertNotIn("private-meta", json.dumps(public))
        self.assertTrue(public["kiwify_client_secret_configured"])
        self.assertEqual(self.store.settings(I)["kiwify_client_secret"], "very-private")

    def test_pending_refunds_and_other_products_never_revenue(self):
        sales = [self.sale(), self.sale("pending", status="waiting_payment"), self.sale("refund", status="refunded"),
                 self.sale("other", product_id="clinic"), self.sale("chargeback", status="chargedback")]
        self.store.save_records(I, C, "kiwify", sales)
        r = self.report()
        self.assertEqual(r["summary"]["gross"], 100000)
        self.assertEqual(r["summary"]["sales"], 1)
        self.assertEqual(r["orders"]["total"], 4)

    def test_first_lead_dedup_and_period_conversion(self):
        self.store.save_records(I, C, "sheets", [self.lead(), self.lead("again", created_day="2026-09-03")])
        self.store.save_records(I, C, "kiwify", [self.sale(), self.sale("repeat")])
        r = self.report()
        self.assertEqual(r["summary"]["leads"], 1)
        self.assertEqual(r["summary"]["conversion"], 100)
        self.assertEqual(r["summary"]["buyers"], 1)

    def test_no_name_only_or_conflicting_phone_match(self):
        lead = self.lead(email="other@example.invalid")
        self.assertEqual(match_lead(self.sale(), [lead])[1], "unmatched")
        self.assertEqual(match_lead(self.sale(email="", phone=""), [lead])[1], "unmatched")

    def test_ambiguous_campaign_not_guessed_and_sale_utm_wins(self):
        leads = [self.lead(), self.lead("b", campaign="different")]
        self.store.save_records(I, C, "sheets", leads)
        self.store.save_records(I, C, "kiwify", [self.sale()])
        self.assertEqual(self.report()["warnings"]["sale_attribution"], 1)
        self.store.save_records(I, C, "kiwify", [self.sale(campaign="direct-utm")])
        r = self.report()
        self.assertEqual(next(c for c in r["campaigns"] if c["sales"])["name"], "direct-utm")
        self.assertNotIn("sale_attribution", r["warnings"])

    def test_meta_mapping_money_and_no_organic_roas(self):
        self.store.save_campaigns(I, C, [{"id":"123", "name":"Curso Meta", "aliases":["co2"]}])
        self.store.save_records(I, C, "sheets", [self.lead()])
        self.store.save_records(I, C, "kiwify", [self.sale(), self.sale("organic", campaign="organic", email="other@example.invalid")])
        self.store.save_records(I, C, "meta", [{"id":"123:2026-09-01", "campaign_id":"123", "day":"2026-09-01", "spend":10000}])
        self.assertIsNone(self.report()["summary"]["spend"])
        self.store.set_sync_state(I, C, "meta", {"ok":True, "from":"2026-09-01", "to":"2026-09-30"})
        r=self.report()
        self.assertEqual(r["summary"]["spend"],10000)
        self.assertEqual(r["summary"]["roas"],10)
        paid=next(c for c in r["campaigns"] if c["meta"])
        self.assertEqual((paid["cpl"],paid["cac"]),(10000,10000))

    def test_missing_campaign_organic_label_preserves_totals_and_source(self):
        self.store.save_records(I, C, "sheets", [self.lead(campaign="")])
        self.store.save_records(I, C, "kiwify", [self.sale()])
        report = self.report()
        bucket = report["campaigns"][0]
        self.assertEqual(bucket["id"], "unattributed")
        self.assertEqual(bucket["name"], "Sem campanha identificada (orgânico)")
        self.assertEqual(report["orders"]["rows"][0]["campaign"], bucket["name"])
        self.assertEqual((bucket["leads"], bucket["sales"], bucket["gross"], bucket["net"]),
                         (1, 1, 100000, 90000))
        self.assertIsNone(bucket["roas"])
        self.assertEqual(self.store.records(I, C, "kiwify")[0]["campaign"], "")

    def test_invalid_values_dates_flagged_not_invented(self):
        self.store.save_records(I,C,"kiwify",[self.sale(approved_day=""),self.sale("foreign",currency="USD")])
        r=self.report()
        self.assertEqual(r["summary"]["sales"],0)
        self.assertTrue(r["totals_partial"])
        self.assertEqual(len(r["pending"]),2)

    def test_mapping_duplicates_and_cross_course_rejected(self):
        with self.assertRaises(ValueError):
            self.store.save_campaigns(I,C,[{"id":"123","name":"A","aliases":["co2"]},{"id":"456","name":"B","aliases":["co2"]}])
        self.store.save_campaigns(I,C,[{"id":"123","name":"A","aliases":[]}])
        self.store.save_course(I,{"key":"other-course","name":"Outro","product_ids":[]})
        with self.assertRaises(ValueError):
            self.store.save_campaigns(I,"other-course",[{"id":"123","name":"A","aliases":[]}])

    def test_mapping_whitespace_cannot_hide_duplicate_aliases(self):
        with self.assertRaises(ValueError):
            self.store.save_campaigns(I,C,[{"id":"123","name":"A","aliases":["co2"]},
                                          {"id":"456","name":"B","aliases":[" co2 "]}])

    def test_account_change_invalidates_only_its_source(self):
        self.store.save_settings(I,{"kiwify_client_id":"original", "kiwify_account_id":"account"})
        course=self.store.course(I,C);course.pop("institute_key");course["product_ids"]=["regen"]
        self.store.save_course(I,course)
        self.store.save_records(I,C,"kiwify",[self.sale()])
        self.store.save_records(I,C,"sheets",[self.lead()])
        self.store.set_sync_state(I,C,"kiwify",{"ok":True,"at":1})
        self.store.save_settings(I,{"kiwify_client_secret":"rotated-secret"})
        self.assertEqual(len(self.store.records(I,C,"kiwify")),1)
        self.store.save_settings(I,{"kiwify_account_id":"another-account"})
        self.assertEqual(self.store.records(I,C,"kiwify"),[])
        self.assertNotIn("kiwify",self.store.sync_state(I,C))
        self.assertEqual(self.store.course(I,C)["product_ids"],[])
        self.assertEqual(len(self.store.records(I,C,"sheets")),1)

    def test_meta_account_change_removes_old_campaign_mapping(self):
        self.store.save_settings(I,{"meta_account_id":"123456"})
        self.store.save_campaigns(I,C,[{"id":"123","name":"A","aliases":["co2"]}])
        self.store.save_records(I,C,"meta",[{"id":"123:2026-09-01","spend":1}])
        self.store.set_sync_state(I,C,"meta",{"ok":True,"at":1})
        self.store.save_settings(I,{"meta_account_id":"999999"})
        self.assertEqual(self.store.campaigns(I,C),[])
        self.assertEqual(self.store.records(I,C,"meta"),[])
        self.assertNotIn("meta",self.store.sync_state(I,C))

    def test_course_source_change_resets_freshness(self):
        self.store.set_sync_state(I,C,"sheets",{"ok":True,"at":1})
        self.store.set_sync_state(I,C,"kommo",{"ok":True,"at":1})
        course=self.store.course(I,C);course.pop("institute_key")
        course.update(sheet_gid="1",pipeline_name="NOVO FUNIL")
        self.store.save_course(I,course)
        self.assertNotIn("sheets",self.store.sync_state(I,C))
        self.assertNotIn("kommo",self.store.sync_state(I,C))

    def test_success_replaces_only_imported_period(self):
        self.store.save_records(I,C,"kiwify",[self.sale(),self.sale("old",created_day="2025-01-01")])
        self.store.save_records(I,C,"kiwify",[],window=("created_day","2026-09-01","2026-09-30"))
        self.assertEqual([r["id"] for r in self.store.records(I,C,"kiwify")],["old"])

    def test_export_all_pages_and_no_formula_injection(self):
        from openpyxl import load_workbook
        self.store.save_records(I,C,"kiwify",[self.sale(str(n),name="=HYPERLINK(1)") for n in range(65)])
        data=workbook(self.store,I,C,{"from":"2026-09-01","to":"2026-09-30"})
        book=load_workbook(io.BytesIO(data))
        self.assertEqual(book["Pedidos"].max_row,66)
        self.assertEqual(book["Pedidos"]["C2"].data_type,"s")

    def test_brazil_approval_date_and_historical_windows(self):
        self.assertEqual(day("2026-10-01T01:00:00Z"),"2026-09-30")
        self.assertEqual(day("30/09/2026 12:35"),"2026-09-30")
        self.assertGreater(len(list(windows("2025-01-01","2026-10-04"))),20)
        self.assertTrue(all((b-a).days<90 for a,b in windows("2025-01-01","2026-10-04")))

    def test_sheet_schema_and_dates(self):
        rows=sheet_rows('Seu nome,Seu Whatsapp,Seu melhor e-mail,Pontuação,Data,ID,utm_campaign\nTeste,11999990000,test@example.invalid,10,30/09/2026,abc,co2\n')
        self.assertEqual(rows[0]["created_day"],"2026-09-30")
        self.assertEqual(rows[0]["score"],"10")

    def test_failed_sources_preserve_previous_data(self):
        self.store.save_records(I,C,"kiwify",[self.sale()])
        with patch("institute_api.Kiwify") as k,patch("institute_api.sheets",side_effect=SourceError("Google indisponível")),patch("institute_api.kommo",side_effect=SourceError("Kommo indisponível")):
            k.return_value.sales.side_effect=SourceError("Kiwify indisponível")
            result=sync_sources(self.store,I,C,"2026-09-01","2026-09-30",lambda path:{},lambda source:None)
        self.assertFalse(result["kiwify"]["ok"])
        self.assertEqual(len(self.store.records(I,C,"kiwify")),1)

    def test_kommo_excludes_clinic_funnels(self):
        def call(path):
            return {"_embedded":{"pipelines":[{"id":1,"name":"CLINICA"},{"id":2,"name":"REGENCODE","_embedded":{"statuses":[]}}]}} if "pipelines" in path else {"_embedded":{"leads":[{"id":1,"pipeline_id":1}]}}
        with self.assertRaises(SourceError):
            kommo(self.store.course(I,C),call)

    def test_kiwify_dedupe_month_overlap_and_full_details(self):
        calls=[]
        def http(url,headers=None,data=None,**kwargs):
            calls.append(url)
            if "oauth" in url:return {"access_token":"private-token"}
            return {"data":[{"id":"abc","product":{"id":"regen"},"status":"paid","approved_date":"2026-09-02T12:00:00Z","created_at":"2026-09-01T12:00:00Z","currency":"BRL","payment":{"charge_amount":12345,"net_amount":11000,"charge_currency":"BRL","settlement_currency":"BRL"}}]}
        source=Kiwify({"kiwify_client_id":"id","kiwify_client_secret":"secret","kiwify_account_id":"account"},http=http,pause=lambda _:None)
        rows=source.sales(self.store.course(I,C),"2026-08-01","2026-09-30")
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]["gross"],12345)
        self.assertTrue(all("view_full_sale_details=true" in url for url in calls if "sales" in url))

    def test_meta_filters_and_currency_guard(self):
        calls=[]
        def http(url,headers=None,**kwargs):
            calls.append((url,headers))
            return {"data":[{"campaign_id":"999","account_currency":"BRL","date_start":"2026-09-01","spend":"10.15"}]}
        meta=Meta({"meta_account_id":"123456","meta_access_token":"secret"},http)
        with self.assertRaises(SourceError):
            meta.insights([{"id":"123"}],"2026-09-01","2026-09-30")
        self.assertIn("filtering",calls[0][0])
        self.assertNotIn("secret",calls[0][0])
        self.assertEqual(calls[0][1]["Authorization"],"Bearer secret")


if __name__ == "__main__":
    unittest.main()
