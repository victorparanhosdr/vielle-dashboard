import io
import json
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from chart_export import CHARTS, build_workbook


class ChartExportTests(unittest.TestCase):
    def test_all_charts_empty_and_formula_safety(self):
        for chart in CHARTS:
            data = build_workbook(chart, {}, {"Clínica": "=HYPERLINK(\"bad\")"})
            workbook = load_workbook(io.BytesIO(data))
            self.assertIn("Contexto", workbook.sheetnames)
            self.assertEqual(workbook["Contexto"]["B3"].data_type, "s")
            self.assertEqual(workbook.worksheets[0].freeze_panes, "A2")

    def test_full_ranking_and_long_source_json(self):
        raw = json.dumps({"text": "a" * 65000})
        sales = [{"uuid": str(i), "patient_uuid": str(i), "patient": "=1+1", "amount": 10,
                  "raw_json": raw if i == 0 else "{}"} for i in range(15)]
        wb = load_workbook(io.BytesIO(build_workbook("top_patients", {"export_details": {"sales": sales}}, {})))
        self.assertEqual(wb["Ranking completo"].max_row, 16)
        self.assertEqual(wb["Vendas"].max_row, 16)
        self.assertEqual(wb["Vendas"]["C2"].data_type, "s")
        original = "".join(row[3] for row in wb["Registros originais"].iter_rows(min_row=2, values_only=True) if row[1] == "0")
        self.assertEqual(original, raw)

    def test_actual_sql_scopes_and_totals(self):
        import app
        with tempfile.TemporaryDirectory() as directory, \
                patch.object(app, "clinic_db_path", return_value=Path(directory) / "clinic.sqlite3"), \
                patch.object(app, "clinic_pipeline_doctor_map", return_value={}), \
                patch("urllib.request.urlopen", side_effect=AssertionError("No external requests")):
            app.init_db()
            with app.db() as conn:
                conn.execute("insert into pipelines(id,name,raw_json,synced_at) values(1,'Funil teste','{}',0)")
                for i, day in enumerate(["2026-09-01", "2026-08-01"]):
                    timestamp = int(datetime.fromisoformat(day + "T12:00:00").timestamp())
                    conn.execute("insert into leads(id,name,pipeline_id,created_at,raw_json,synced_at) values(?,?,1,?,'{}',0)",
                                 (i + 1, f"Lead teste {i}", timestamp))
                for i, day in enumerate(["2026-09-01", "2026-08-01"]):
                    conn.execute("insert into clinica_bills(uuid,type,due_date,amount,raw_json,synced_at) values(?,?,?,?,?,0)",
                                 (f"bill{i}", "sale", day, 999, '{"final_amount":12345}'))
                conn.execute("insert into clinica_parcels(uuid,type,status,due_date,amount,category_name,raw_json,synced_at) values('extra','other','received','2026-09-01',20,'Manual','{}',0)")
                conn.execute("insert into clinica_parcels(uuid,type,status,due_date,amount,category_name,raw_json,synced_at) values('expense','expense','paid','2026-09-01',999,'Aluguel','{\"final_amount\":3000}',0)")
                for i in range(12):
                    raw = json.dumps({"status": "active", "final_amount": 10000, "buyer": {"name": f"Teste {i}"}})
                    conn.execute("insert into clinica_sales(uuid,patient_uuid,sale_date,total,raw_json,synced_at) values(?,?,?,?,?,0)",
                                 (f"sale{i}", f"patient{i}", "2026-09-01", 100, raw))
                conn.execute("insert into clinica_bookings(uuid,registered_at,starts_at,raw_json,synced_at) values('booking','2026-09-01','2026-10-01','{}',0)")
            args = {"date_from": "2026-09-01", "date_to": "2026-09-02"}
            regular = app.report_data(**args)["general_panel"]
            panel = app.report_data(**args, export_chart="leads_bookings")["general_panel"]
            self.assertNotIn("export_details", regular)
            details = panel.pop("export_details")
            self.assertEqual(regular, panel)
            self.assertAlmostEqual(sum(r["amount"] for r in details["income"]), 143.45)
            self.assertAlmostEqual(sum(r["income"] for r in panel["financial_daily"]), 143.45)
            self.assertAlmostEqual(sum(r["amount"] for r in details["expenses"]), 30)
            self.assertEqual(len(details["sales"]), 12)
            self.assertEqual(len(panel["top_patients"]), 10)
            self.assertEqual(details["bookings"][0]["day"], "2026-09-01")
            self.assertEqual(len(details["leads"]), 1)
            self.assertEqual(details["leads"][0]["id"], 1)
            self.assertEqual(sum(r["total"] for r in panel["daily_leads"]), 1)
            self.assertEqual(sum(r["total"] for r in panel["daily_bookings"]), 1)
            panel["export_details"] = details
            for chart in CHARTS:
                load_workbook(io.BytesIO(build_workbook(chart, panel, {})))


if __name__ == "__main__":
    unittest.main()
